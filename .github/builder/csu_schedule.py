"""
Парсер расписания ЧелГУ — сетка Excel, одинаковая на всех факультетах:
листы по потокам, колонка с днями недели, колонка с номером пары,
у каждой группы свои столбцы (обычно 2 — это подгруппы).

Каждая пара занимает блок из 4 строк:
  строки 1–2 — нечётная неделя, строки 3–4 — чётная,
  запись по центру блока — каждую неделю.
Префиксы недель: «1н» — нечётные, «2н» — чётные, «2,4,6,8н» / «1-4н» — конкретные недели.
Проверено на всех файлах сайта (сентябрь 2026): «1н» стоит в верхней (нечётной) половине
блока ~2000 раз, «2н» — в нижней (чётной) ~1850 раз, обратных случаев единицы.
Бывает и запись в скобках: «(2-8Н.) ИНОСТРАННЫЙ ЯЗЫК», «ИСТОРИЯ РОССИИ (ПР.) (1-8Н.)».
"""

import hashlib
import html
import io
import re
from collections import OrderedDict
from urllib.parse import urljoin
from datetime import time

from openpyxl import load_workbook

# Расписание звонков ФЭУ ЧелГУ (econom.csu.ru/education/call-schedule/)
PAIR_TIMES = {
    1: time(8, 0), 2: time(9, 40), 3: time(11, 20), 4: time(13, 15),
    5: time(15, 0), 6: time(16, 40), 7: time(18, 20), 8: time(19, 55),
}

DAYS = {"ПОНЕДЕЛЬНИК": 0, "ВТОРНИК": 1, "СРЕДА": 2, "ЧЕТВЕРГ": 3, "ПЯТНИЦА": 4, "СУББОТА": 5}

_WEEKS = r"((?:\d+\s*[,\-–]\s*)*\d+)\s*[НнHh]"
WEEK_PREFIX = re.compile(r"^\s*" + _WEEKS + r"(?:\.\s*|\s+)", re.IGNORECASE)   # «1н ПРЕДМЕТ», «2,4н. ПРЕДМЕТ»
# «(1-8Н.)» или «С 1-5НЕД» в любом месте строки
WEEK_PAREN = re.compile(r"(?:\(\s*|\bС\s+)" + _WEEKS + r"(?:ЕД)?\.?(?=[\s),]|$)\s*\)?", re.IGNORECASE)
URL = re.compile(r"\s*https?://[^\s)]+")
# Строка преподавателя: инициалы («ИВАНОВ И.И.», «ТКАЧЕНКО В.,»), «АУД.», должность или «218(2)» в конце
TEACHER_LINE = re.compile(
    r"[А-ЯЁ]\s?\.\s*[А-ЯЁ]\.|[А-ЯЁ]{2,}\s+[А-ЯЁ]\.\s*,|\bАУД\.|\(\d\)\s*$|\(\d\)\s*/"
    r"|,\s*(?:ДОЦ|ПРЕП|СТ\.\s*ПРЕП|ПРОФ|ЗАВ)|,\s*(?:[А-ЯA-Z]-)?\d{2,4}\s?[А-ЯA-Z]?$", re.IGNORECASE)
ROOM = re.compile(r"([А-ЯA-Z]?-?\d+[А-Яа-я]?|Л/З)\((\d)\)")   # формат ФЭУ: «218(2)», «Л/З(2)»
AUD = re.compile(r"\bАУД\.?\s*", re.IGNORECASE)                 # «…, ДОЦЕНТ, АУД. А-23(1К)»
TAIL_ROOM = re.compile(                                          # «…, ДОЦ., 117 Б (4 КОРПУС)», «…, А-15»
    r"(?:^|[,\s])((?:[А-ЯA-Z]-)?\d{2,4}\s?[А-ЯA-Z]?(?:\s*/\s*[\w.]+)*(?:\s*\([^)]*\)?)?)\s*$")
BUILDING = re.compile(r",?\s*\((\d+)\s*К(?:ОРП)?\.?\)\s*$", re.IGNORECASE)  # «(11КОРП)», «(4КОРП.)»
SKIP = re.compile(r"САМОСТОЯТ|^ДЕНЬ$|^ЗАНЯТИЙ$|^РАБОТЫ$", re.IGNORECASE)
# Пометки потока над названием пары: «МН-1», «21М,Э-2» — не предмет;
# и совместные группы: «+МН-3», «+ЛПА-401, ЛМБА-401», «САб-11»
TAG = re.compile(r"^\+|^(?=.*\d)[А-ЯЁа-яёA-Za-z\d,\-.]{1,12}$")
# «до 30.10» / «с 01.10» над названием пары — пара идёт только до / с этой даты
DATE_NOTE = re.compile(r"^(до|по|с|со)\s+(\d{1,2})\.(\d{1,2})\.?$", re.IGNORECASE)
ONLINE = re.compile(r"ОНЛАЙН|ДИСТАНЦ", re.IGNORECASE)


FAC_PREFIX = re.compile(r"^\s*[А-ЯЁA-Z]{1,3}\d-")


def norm_group(s: str) -> str:
    """«ИО1-СДОб-11(16)» -> «ИО1СДОБ11»"""
    return re.sub(r"[\s\-_]", "", re.sub(r"\(.*?\)", "", str(s))).upper()


def group_matches(header: str, code: str) -> bool:
    """Совпадение кода группы с шапкой столбца; префикс факультета («ЭУ1-») можно не писать.
    В шапке бывает несколько групп через запятую — подходит любая из них или вся шапка целиком."""
    if norm_group(header) == norm_group(code):
        return True
    full, short = norm_group(code), norm_group(FAC_PREFIX.sub("", code))
    for part in re.split(r"[,;]", str(header)):
        if norm_group(part) == full or norm_group(FAC_PREFIX.sub("", part.strip())) == short:
            return True
    return False


def discover_links(page_html: str, base_url: str) -> list[tuple[str, str]]:
    """Ссылки на .xlsx со страницы расписаний: [(название факультета, url), ...]"""
    hrefs = list(re.finditer(r'href="([^"]+\.xlsx)"', page_html, re.I))
    out = []
    for i, m in enumerate(hrefs):
        end = hrefs[i + 1].start() if i + 1 < len(hrefs) else m.end() + 3000
        t = html.unescape(re.sub(r"<[^>]+>", " ", page_html[m.end():end])).replace("​", "")
        t = re.sub(r"\s+", " ", t)
        t = re.split(r"[\d.,]+\s*(?:Кбайт|Мбайт|КБ|МБ|KB|MB)", t)[0].strip(" >")
        out.append((t, urljoin(base_url, m.group(1))))
    return out


def title_key(title: str) -> str:
    return re.sub(r"\W", "", title or "").lower()


def list_groups(data: bytes) -> list[str]:
    """Быстро читает только шапку каждого листа (read_only)."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True)
    except Exception:
        return []
    result = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(max_row=60, values_only=True)]
        for i, row in enumerate(rows):
            day_idx = next((j for j, v in enumerate(row)
                            if isinstance(v, str) and v.strip().upper() == "ПОНЕДЕЛЬНИК"), None)
            if day_idx is not None and i > 0:
                result += [str(v).strip() for j, v in enumerate(rows[i - 1])
                           if j > day_idx + 1 and v and str(v).strip()]
                break
    wb.close()
    return result


def _find_layout(ws):
    """-> (day_col, pair_col, first_row, [(group_name, col_from, col_to), ...])"""
    for row in ws.iter_rows(max_row=60):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().upper() == "ПОНЕДЕЛЬНИК":
                day_col, pair_col, first = c.column, c.column + 1, c.row
                header = [(x.column, str(x.value).strip()) for x in ws[first - 1]
                          if x.value and x.column > pair_col and str(x.value).strip()]
                groups = []
                for i, (col, name) in enumerate(header):
                    nxt = header[i + 1][0] - 1 if i + 1 < len(header) else col + 1
                    groups.append((name, col, nxt))
                return day_col, pair_col, first, groups
    return None


def _week_numbers(spec: str) -> set[int]:
    nums = set()
    for part in re.split(r"\s*,\s*", spec.strip()):
        if re.search(r"[\-–]", part):
            a, b = [int(x) for x in re.split(r"\s*[\-–]\s*", part)]
            nums.update(range(a, b + 1))
        elif part:
            nums.add(int(part))
    return nums


def _week_rule(prefix: str | None, half: str) -> str:
    if prefix:
        nums = _week_numbers(prefix)
        if nums == {1}:
            return "нечет"
        if nums == {2}:
            return "чет"
        # «(1-8Н.)» в верхней половине блока — нечётные недели из 1–8, в нижней — чётные
        parity = {"top": 1, "bottom": 0}.get(half)
        if parity is not None:
            nums = {n for n in nums if n % 2 == parity} or nums
        return ",".join(str(n) for n in sorted(nums))
    return {"top": "нечет", "bottom": "чет"}.get(half, "все")


def split_subject(text: str) -> tuple[str | None, str]:
    """-> (недели или None, название предмета без пометки недель)"""
    m = WEEK_PREFIX.match(text)
    if m:
        return m.group(1), text[m.end():].strip()
    m = WEEK_PAREN.search(text)
    if m:
        return m.group(1), (text[:m.start()] + " " + text[m.end():]).strip()
    return None, text


def _parse_teacher(text: str) -> tuple[str, str]:
    """Строка преподавателя -> (преподаватель, аудитория)"""
    text = WEEK_PAREN.sub("", re.sub(r"\s+", " ", text)).strip()
    if "http" in text:  # дистанционная пара со ссылкой
        return text.split("http")[0].strip(" /,"), "онлайн"
    m = re.search(r",?\s*\(?(?:ОНЛАЙН|ДИСТАНЦИОННО)\)?\s*$", text, re.IGNORECASE)
    if m:  # «…, ДОЦЕНТ, ОНЛАЙН»
        return text[:m.start()].strip(" /,"), "онлайн"
    if ROOM.search(text):  # формат ФЭУ: «ИВАНОВ И.И.(доцент) 218(2) / 219(2)»
        rooms = [f"{r} (корп. {k})" for r, k in ROOM.findall(text)]
        return text[:ROOM.search(text).start()].strip(" /,"), " / ".join(rooms)
    m = AUD.search(text)
    if m:  # «…, ДОЦЕНТ, АУД. А-23(1К)»
        return text[:m.start()].strip(" /,"), text[m.end():].strip(" ,")
    m = re.search(r",\s*((?:УЛ\.|ГАУЗ|ГБУЗ|МАУЗ).*)$", text, re.IGNORECASE)
    if m:  # занятие по адресу: «…, УЛ.КУЗНЕЦОВА 2А», «…, ГАУЗ ДГКП № 8 …»
        return text[:m.start()].strip(" /,"), m.group(1)
    m = re.search(r",\s*(ФОК|БАССЕЙН|СПОРТЗАЛ|С/З|\d+\s*К(?:ОРП)?\.?)\s*$", text, re.IGNORECASE)
    if m:  # «…, ПРЕП., ФОК», «…, ДЕК.ФАК., 3К.»
        return text[:m.start()].strip(" /,"), m.group(1)
    m = BUILDING.search(text)
    if m:  # «…, ПРЕПОДАВАТЕЛЬ, (11КОРП)»
        return text[:m.start()].strip(" /,"), f"корп. {m.group(1)}"
    m = TAIL_ROOM.search(text)
    if m and not re.search(r"[А-ЯЁ]\.$", text):  # «…, ДОЦ., 117 Б (4 КОРПУС)», «…, А-15»
        return text[:m.start(1)].strip(" /,"), m.group(1).strip()
    return text.strip(" /,"), ""


# Разобранные книги: sha1(файл) -> {шапка группы: [занятия с номером пары]}.
# Открытие книги занимает ~3 с, поэтому книга разбирается один раз на все её группы.
_BOOKS: "OrderedDict[str, dict[str, list[dict]]]" = OrderedDict()
_BOOKS_MAX = 20


def _parse_sheet_group(ws, cell_range, day_rows, first_row, pair_col, c1, c2) -> list[dict]:
    width = c2 - c1 + 1
    rows = []
    for r in range(first_row, ws.max_row + 1):
        v = ws.cell(r, pair_col).value
        if v is None or not str(v).strip().isdigit():
            continue
        pair_no = int(str(v).strip())
        weekday = next((wd for dr, wd in reversed(day_rows) if dr <= r), None)
        if weekday is None:
            continue
        rng = cell_range.get((r, pair_col))
        b_start, b_end = (rng[0], rng[2]) if rng else (r, r + 3)

        # собираем все тексты блока в столбцах группы
        seen, entries = set(), []
        for rr in range(b_start, b_end + 1):
            for cc in range(c1, c2 + 1):
                box = cell_range.get((rr, cc), (rr, cc, rr, cc))
                if box in seen:
                    continue
                seen.add(box)
                text = ws.cell(box[0], box[1]).value
                if not isinstance(text, str) or not text.strip():
                    continue
                col_from, col_to = max(box[1], c1), min(box[3], c2)
                sub = "" if (width == 1 or (col_from == c1 and col_to == c2)) else str(col_from - c1 + 1)
                entries.append((box[0], box[2], sub, text.strip()))

        # разбираем по подгруппам: строка предмета + строка преподавателя
        for sub in sorted({e[2] for e in entries}):
            items = sorted((e for e in entries if e[2] == sub), key=lambda e: e[0])
            current, note = None, {}

            def finish():
                offset, last = current.pop("_top") - b_start, current.pop("_bottom") - b_start
                if offset == 1 or (offset == 0 and last >= 2):
                    half = "all"       # запись по центру блока или на весь блок
                elif offset == 0:
                    half = "top"
                else:
                    half = "bottom"
                current["week"] = _week_rule(current.pop("_prefix"), half)
                subject = current["subject"]
                if not current.get("link") and (lm := URL.search(subject)):
                    current["link"] = lm.group(0).strip()
                current["subject"] = re.sub(r"\s+", " ", URL.sub("", subject)).strip()
                if not current["room"] and (ONLINE.search(subject) or "http" in subject):
                    current["room"] = "онлайн"
                if current["subject"]:
                    rows.append(current)

            for top, bottom, _, text in items:
                text = re.sub(r"\s+", " ", text).strip()
                if current and text.startswith("http"):  # ссылка на дистанционную пару
                    current["room"] = "онлайн"
                    current["link"] = URL.search(text).group(0).strip()
                    continue
                if TEACHER_LINE.search(text) and current:
                    if lm := URL.search(text):
                        current["link"] = lm.group(0).strip()
                    current["teacher"], current["room"] = _parse_teacher(text)
                    if current["_prefix"] is None and (m := WEEK_PAREN.search(text)):
                        current["_prefix"] = m.group(1)  # «ДАНИЛОВ А.Р., СТ.ПРЕП. (9-17Н.)»
                    finish()
                    current = None
                    continue
                prefix, subject = split_subject(text)
                # название, разбитое на несколько строк: «НАУЧНО-» / «ИССЛЕДОВАТЕЛЬСКАЯ» / «РАБОТА»
                if (current and prefix is None and top == current["_bottom"] + 1
                        and not (TAG.match(text) or DATE_NOTE.match(text) or SKIP.search(text))
                        and not current["subject"].endswith(")")):
                    joiner = "" if current["subject"].endswith("-") else " "
                    current["subject"] += joiner + subject
                    current["_bottom"] = bottom
                    continue
                if current:
                    finish()
                    current = None
                if m := DATE_NOTE.match(text):
                    note["until" if m.group(1).lower() in ("до", "по") else "since"] = \
                        f"{int(m.group(3)):02d}-{int(m.group(2)):02d}"
                    continue
                if SKIP.search(text) or TAG.match(text) or TEACHER_LINE.search(text) or text.startswith("http"):
                    continue  # «день самостоятельной работы», пометка потока, преподаватель без предмета
                current = {
                    "date": None, "weekday": weekday, "pair": pair_no, "subject": subject,
                    "room": "", "teacher": "", "subgroup": sub, "link": "",
                    "_top": top, "_bottom": bottom, "_prefix": prefix, **note,
                }
                note = {}
            if current:
                finish()
    return rows


def parse_book(data: bytes) -> dict[str, list[dict]]:
    """Все группы книги за один проход -> {шапка группы: [занятия]}; результат кэшируется."""
    key = hashlib.sha1(data).hexdigest()
    if key in _BOOKS:
        _BOOKS.move_to_end(key)
        return _BOOKS[key]
    wb = load_workbook(io.BytesIO(data))
    book: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        layout = _find_layout(ws)
        if not layout:
            continue
        day_col, pair_col, first_row, groups = layout
        # карта объединённых ячеек: каждая клетка -> диапазон
        cell_range = {}
        for rng in ws.merged_cells.ranges:
            box = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    cell_range[(r, c)] = box
        day_rows = []
        for r in range(first_row, ws.max_row + 1):
            v = ws.cell(r, day_col).value
            if isinstance(v, str) and v.strip().upper() in DAYS:
                day_rows.append((r, DAYS[v.strip().upper()]))
        for name, c1, c2 in groups:
            if name not in book:
                book[name] = _parse_sheet_group(ws, cell_range, day_rows, first_row, pair_col, c1, c2)
    wb.close()
    _BOOKS[key] = book
    while len(_BOOKS) > _BOOKS_MAX:
        _BOOKS.popitem(last=False)
    return book


def parse_csu(data: bytes, group: str, pair_times: dict | None = None) -> list[dict]:
    times = pair_times or PAIR_TIMES
    book = parse_book(data)
    name = next((n for n in book if group_matches(n, group)), None)
    if name is None:
        raise ValueError(f"Группа «{group}» не найдена в расписании")
    return [{k: v for k, v in r.items() if k != "pair"} | {"start": times[r["pair"]]}
            for r in book[name] if r["pair"] in times]


def teacher_rows(data: bytes, surname: str, pair_times: dict | None = None) -> list[tuple[str, dict]]:
    """Занятия, где в строке преподавателя встречается фамилия -> [(шапка группы, строка как в parse_csu)]."""
    times = pair_times or PAIR_TIMES
    q = surname.upper().replace("Ё", "Е")
    return [(name, {k: v for k, v in r.items() if k != "pair"} | {"start": times[r["pair"]]})
            for name, rows in parse_book(data).items() for r in rows
            if r["pair"] in times and q in r["teacher"].upper().replace("Ё", "Е")]


# ───── Оформление для показа (общее для бота и приложения tools/build_app_data.py) ─────

SMALL_WORDS = {"в", "во", "и", "на", "по", "для", "с", "со", "к", "о", "об", "от", "из", "за", "не", "или", "а", "но", "у",
               "при", "над", "под", "про", "до"}


PROPER_WORDS = {"россия", "россии", "российской", "российский", "федерации", "урала", "уральский", "челябинской"}


def pretty_subject(subject: str) -> str:
    """Предметы в файлах ЧелГУ написаны КАПСОМ — показываем нормальным регистром; короткие аббревиатуры (ИТ, ООП) остаются."""
    bare = re.sub(r"\(.*?\)", "", subject)
    if bare.upper() != bare:
        return subject   # уже набрано нормально
    out = []
    for w in subject.split(" "):
        letters = re.sub(r"[^A-Za-zА-Яа-яЁё]", "", w)
        if not letters or w.startswith("("):
            out.append(w)
        elif letters.lower() in SMALL_WORDS and out:
            out.append(w.lower())
        elif len(letters) <= 3:
            out.append(w)
        elif letters.lower() in PROPER_WORDS:
            out.append(w.lower().capitalize())
        else:
            out.append(w.lower())
    text = " ".join(out)
    return text[:1].upper() + text[1:]


def pretty_teacher(name: str) -> str:
    """«МОЛОДЦОВ А.Л.» -> «Молодцов А.Л.»: фамилия с заглавной, инициалы как есть."""
    words = name.split(" ")
    if words and words[0].isupper():
        words[0] = words[0].title()
    return " ".join(words)


def teacher_short(t: str) -> str:
    """«МАВРИНА Н.А.(ст.преп.), ...» -> «МАВРИНА Н.А.»: без должности, чтобы один человек не считался двумя."""
    return re.sub(r"\s*\(.*?\)", "", t.split(",")[0]).strip()
